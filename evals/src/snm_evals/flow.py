"""``snm-sim flow``: play the whole case flow and check its invariants (decision record 005).

The driver plays every side through the same ``CaseService`` calls the pages use: purchasing drafts the
request (the deterministic request text, or the planner agent with ``--planner``), approves it with the
engine's mandate and the request sheet; a scripted supplier downloads the sheet, fills it according to a
behavior (inside the mandate, late, two options, a decline, a message first and the sheet later), and
uploads it; purchasing decides (the recommended proposal, else the engine's lever); the system hands the
decision off and posts the receipt the data would show. Transfer cases are published and received.

After each case the invariants are checked on what the store holds: append-only sequences, legal
transitions and actors, every decision by a person, no figure outside the outward fields in outbound text,
nothing internal on the supplier view, every file stored with its hash and parsed, every proposal under a
mandate flagged and the flags reproducible. The report lists every case with its decision and its actor.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import time
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from snm_cases import excel, model
from snm_cases.planning import LocalPlanningData
from snm_cases.server import REPO_ROOT, default_planning_dir
from snm_cases.service import XLSX, Actor, CaseError, CaseService
from snm_cases.store import LocalCaseStore

BEHAVIORS = ("inside", "late", "two_options", "decline", "message_first")
BUYER = Actor("human", "purchasing")
INTERNAL_KEYS = (
    "surcharge_pct_max",
    "freight_max_usd",
    "within_mandate",
    "value_usd",
    "rationale",
    "triage",
    "lever_cost_usd",
    "stockout_cost_usd",
    "basis",
    "max_exchanges",
)
COST_PARAMETER = REPO_ROOT / "out" / "small" / "raw_ext" / "cost_parameter.parquet"


@dataclass
class CaseRun:
    case_id: str
    kind: str  # order | transfer
    behavior: str
    decision: str | None = None
    decided_by: str | None = None
    recommended: str | None = None
    final_state: str = ""
    events: int = 0
    proposals: int = 0
    in_mandate: int = 0
    attachments: int = 0
    seconds: float = 0.0
    violations: list[str] = field(default_factory=list)
    error: str | None = None


# ----------------------------------------------------------------------------------------------- supplier


def fill_sheet(sheet: bytes, rows: list[dict[str, Any]]) -> bytes:
    book = load_workbook(io.BytesIO(sheet))
    table = book[excel.SHEET]
    keys = [key for key, _ in excel.COLUMNS]
    for index, row in enumerate(rows, start=1):
        for column, key in enumerate(keys, start=1):
            if key in row:
                table.cell(row=excel.TABLE_ROW + index, column=column, value=row[key])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def supplier_rows(behavior: str, mandate: dict[str, Any], price: float, today: date) -> list[dict[str, Any]]:
    """What the scripted supplier offers: priced on its own price, dated around the needed date."""
    latest = date.fromisoformat(str(mandate["latest_arrival_date"]))
    kind = next(k for k in mandate["allowed_kinds"] if k != "partial")
    quantity = float(mandate["quantity_max"])
    valid = (today + timedelta(days=int(mandate["min_validity_days"]) + 2)).isoformat()
    surcharge = round(float(mandate["surcharge_pct_max"]) * 0.5, 1)
    freight = round(float(mandate["freight_max_usd"]) * 0.8, 2)
    base = {"kind": kind, "quantity": quantity, "unit_price": round(price, 2), "valid_until": valid}
    inside = base | {
        "arrival_date": (latest - timedelta(days=1)).isoformat(),
        "ship_date": (latest - timedelta(days=3)).isoformat(),
    }
    if surcharge > 0:
        inside["surcharge_pct"] = surcharge
    if freight > 0:
        inside["freight_amount"] = freight
    if behavior in ("inside", "message_first"):
        return [inside]
    if behavior == "late":
        return [inside | {"arrival_date": (latest + timedelta(days=5)).isoformat()}]
    if behavior == "two_options":
        half = max(1.0, float(int(quantity // 2)))
        split = (
            f"{half:g}@{(latest - timedelta(days=2)).isoformat()}; "
            f"{quantity - half:g}@{(latest + timedelta(days=10)).isoformat()}"
        )
        dearer = inside | {"unit_price": round(price * 1.1, 2)}
        partial = {
            "kind": "partial",
            "split_deliveries": split,
            "unit_price": round(price, 2),
            "valid_until": valid,
        }
        return [dearer, partial]
    return [{"kind": "decline", "conditions": "Cannot meet the date this month."}]


# ----------------------------------------------------------------------------------------------- flow


def receipt_of(service: CaseService, case_id: str, choice: str) -> tuple[float, date]:
    """The receipt the data would show for the chosen option: its quantity on its (last) arrival date."""
    view = service.view(case_id)
    snap = view.snapshot
    if choice.startswith("P-"):
        proposal = next(p for p in service.store.proposals(case_id) if p.proposal_id == choice)
        if proposal.split_deliveries:
            splits = json.loads(proposal.split_deliveries)
            return sum(float(s["quantity"]) for s in splits), date.fromisoformat(splits[-1]["arrival_date"])
        return float(proposal.quantity or 0.0), proposal.arrival_date or date.fromisoformat(
            str(snap["trigger_date"])
        )
    levers = {snap.get("lever"): {"qty": snap.get("lever_qty"), "coverage_date": snap.get("coverage_date")}}
    for alternative in snap.get("alternatives") or []:
        levers.setdefault(alternative["lever"], alternative)
    lever = levers[choice]
    return float(lever["qty"]), date.fromisoformat(str(lever["coverage_date"] or snap["trigger_date"]))


def play_order(
    service: CaseService, case_id: str, behavior: str, today: date, planner: Any = None
) -> CaseRun:
    run = CaseRun(case_id, "order", behavior)
    snap = service.view(case_id).snapshot
    request_type = "expedite" if snap.get("lever") == "expedite" else "quote"
    if planner is not None:
        planner(case_id)
    if service.view(case_id).state == "new":
        service.draft_supplier_request(
            BUYER, case_id, request_type, service.default_request_text(case_id, request_type)
        )
    view = service.view(case_id)
    draft = next(e for e in reversed(view.events) if e.event_type == "request_drafted")
    defaults = service.mandate_defaults(BUYER, case_id)
    sheet = excel.request_sheet(draft.payload["outward"], draft.payload["message"])
    service.send_request(
        BUYER,
        case_id,
        defaults["mandate"],
        "engine_default",
        defaults["basis"],
        (f"request-{case_id}.xlsx", sheet),
    )
    vendor = str(snap.get("vendor"))
    supplier_view = service.supplier_view(vendor, case_id)
    outbound = next(a for a in supplier_view["attachments"] if a["direction"] == "outbound")
    _, downloaded = service.supplier_file(vendor, case_id, outbound["attachment_id"])
    price = float(snap.get("value_usd") or 0.0) / max(float(snap.get("gap_qty") or 1.0), 1e-9)
    rows = supplier_rows(behavior, defaults["mandate"], price, today)
    filled = fill_sheet(downloaded, rows)
    options, problems = excel.parse_reply(filled, case_id)
    if problems:
        raise RuntimeError(f"the scripted sheet did not parse: {problems}")
    if behavior == "message_first":
        service.supplier_reply(
            vendor, case_id, "We are checking with the factory and will send the sheet.", []
        )
    service.supplier_reply(
        vendor, case_id, "Our answer is in the sheet.", options, (f"reply-{case_id}.xlsx", filled, XLSX)
    )
    run.recommended = service.recommended_proposal(case_id)
    choices = service.decision_choices(case_id)
    if run.recommended:
        choice, rationale = run.recommended, "Purchasing's first-ranked proposal."
    else:
        lever = str(snap.get("lever"))
        fallback = [c for c in choices if not c.startswith("P-") and c != "no_action"]
        if behavior == "decline" and lever == "expedite":
            fallback = [c for c in fallback if c != "expedite"] or fallback
        choice = fallback[0] if fallback else "no_action"
        rationale = "No proposal inside the mandate; the engine's own option instead."
    events = service.decide(BUYER, case_id, choice, rationale, run_id="sim-flow")
    run.decision, run.decided_by = choice, events[0].actor_kind
    if choice != "no_action":
        quantity, received = receipt_of(service, case_id, choice)
        service.confirm_receipt("sim-flow", case_id, quantity, received)
    return run


def play_transfer(service: CaseService, case_id: str) -> CaseRun:
    run = CaseRun(case_id, "transfer", "publish")
    service.publish_transfer("sim-flow", case_id)
    snap = service.view(case_id).snapshot
    service.confirm_receipt(
        "sim-flow",
        case_id,
        float(snap["lever_qty"]),
        date.fromisoformat(str(snap.get("coverage_date") or snap["trigger_date"])),
    )
    run.decision, run.decided_by = str(snap.get("lever")), "system"
    return run


# ----------------------------------------------------------------------------------------------- invariants


def check_case(service: CaseService, case_id: str, earlier: list[Any] | None = None) -> list[str]:
    problems: list[str] = []
    events = service.store.events(case_id)
    if [e.seq for e in events] != list(range(1, len(events) + 1)):
        problems.append("sequence numbers are not 1..n")
    if earlier is not None and events[: len(earlier)] != earlier:
        problems.append("an earlier event changed (append-only broken)")
    state = None
    for event in events:
        transition = model.TRANSITIONS[event.event_type]
        if event.actor_kind not in transition.actors:
            problems.append(f"{event.seq} {event.event_type}: actor {event.actor_kind} not allowed")
        if state is not None and state not in transition.from_states:
            problems.append(f"{event.seq} {event.event_type}: not allowed from state {state}")
        expected = transition.to_state or state
        if event.state_after != expected:
            problems.append(f"{event.seq} {event.event_type}: state {event.state_after}, expected {expected}")
        if event.event_type == "decision_recorded" and event.actor_kind != "human":
            problems.append(f"{event.seq}: a decision recorded by {event.actor_kind}")
        state = event.state_after
    view = service.view(case_id)
    draft = next((e for e in reversed(view.events) if e.event_type == "request_drafted"), None)
    if draft is not None:
        outward = model.OutwardFields(**draft.payload["outward"])
        for message in service.store.messages(case_id):
            if message.direction != "outbound":
                continue
            try:
                service.check_outgoing_text(view, message.body, outward)
            except CaseError as error:
                problems.append(f"{message.message_id}: {error}")
        vendor = str(view.snapshot.get("vendor"))
        if any(e.event_type == "request_sent" for e in view.events):
            shown = json.dumps(service.supplier_view(vendor, case_id))
            leaked = [key for key in INTERNAL_KEYS if key in shown]
            if leaked:
                problems.append(f"supplier view shows {', '.join(leaked)}")
    proposals = service.store.proposals(case_id)
    for attachment in service.store.attachments(case_id):
        content = service.store.read_file(attachment.storage_path)
        if hashlib.sha256(content).hexdigest() != attachment.sha256 or len(content) != attachment.size_bytes:
            problems.append(f"{attachment.attachment_id}: stored file does not match its hash or size")
        if attachment.direction == "inbound":
            parsed = [p for p in proposals if p.source_message_id == attachment.message_id]
            if attachment.parsed_options != len(parsed):
                read, recorded = attachment.parsed_options, len(parsed)
                problems.append(f"{attachment.attachment_id}: {read} options read, {recorded} recorded")
    mandate = service.mandate_of(view)
    for proposal in proposals:
        if mandate is None:
            continue
        if proposal.within_mandate is None:
            problems.append(f"{proposal.proposal_id}: no mandate flags")
            continue
        event = next(
            e
            for e in view.events
            if e.event_type == "proposal_recorded" and proposal.proposal_id in e.payload["proposal_ids"]
        )
        option = model.SupplierProposal.model_validate(
            {
                k: v
                for k, v in event.payload["options"][
                    event.payload["proposal_ids"].index(proposal.proposal_id)
                ].items()
                if k != "within_mandate"
            }
        )
        line_value = service._price_usd(view.snapshot) * mandate.quantity_max
        again = service.within_mandate(mandate, option, line_value, event.created_at.date())
        if again != json.loads(proposal.within_mandate):
            problems.append(f"{proposal.proposal_id}: mandate flags do not reproduce")
    return problems


# ----------------------------------------------------------------------------------------------- run


def sample(service: CaseService, orders: int, transfers: int) -> tuple[list[str], list[str]]:
    # Only cases the latest engine run still flags (a cleared case from an earlier run stays open until a
    # person closes it, and its snapshot may predate the current engine).
    views = sorted((v for v in service.all_views() if not v.cleared), key=lambda v: v.case_id)
    order_ids = [
        v.case_id
        for v in views
        if v.state == "new"
        and v.snapshot.get("condition") in ("shortage", "warning")
        and v.snapshot.get("vendor")
    ]
    transfer_ids = [
        v.case_id for v in views if v.state == "new" and v.snapshot.get("condition") in ("pull", "excess")
    ]
    # Spread the sample over the levers: alternate expedite-type and quote-type cases.
    levers = {v.case_id: v.snapshot.get("lever") for v in views}
    expedite = [c for c in order_ids if levers[c] == "expedite"]
    quote = [c for c in order_ids if c not in expedite]
    picked = []
    while len(picked) < orders and (expedite or quote):
        for pool in (expedite, quote):
            if pool and len(picked) < orders:
                picked.append(pool.pop(0))
    return picked, transfer_ids[:transfers]


def run_flow(
    service: CaseService, orders: int, transfers: int, today: date, planner: Any = None
) -> list[CaseRun]:
    order_ids, transfer_ids = sample(service, orders, transfers)
    runs = []
    for index, case_id in enumerate(order_ids):
        started = time.time()
        earlier = list(service.store.events(case_id))
        try:
            run = play_order(service, case_id, BEHAVIORS[index % len(BEHAVIORS)], today, planner)
        except (CaseError, RuntimeError, StopIteration) as error:
            run = CaseRun(
                case_id, "order", BEHAVIORS[index % len(BEHAVIORS)], error=f"{type(error).__name__}: {error}"
            )
        runs.append(_finish(service, run, earlier, started))
    for case_id in transfer_ids:
        started = time.time()
        earlier = list(service.store.events(case_id))
        try:
            run = play_transfer(service, case_id)
        except CaseError as error:
            run = CaseRun(case_id, "transfer", "publish", error=f"CaseError: {error}")
        runs.append(_finish(service, run, earlier, started))
    return runs


def _finish(service: CaseService, run: CaseRun, earlier: list[Any], started: float) -> CaseRun:
    run.violations = check_case(service, run.case_id, earlier)
    view = service.view(run.case_id)
    run.final_state = view.state
    run.events = len(view.events)
    proposals = service.store.proposals(run.case_id)
    run.proposals = len(proposals)
    run.in_mandate = sum(1 for p in proposals if p.within_mandate and json.loads(p.within_mandate).get("all"))
    run.attachments = len(service.store.attachments(run.case_id))
    run.seconds = round(time.time() - started, 2)
    return run


INVARIANTS = (
    "Invariants checked after each case: sequence numbers 1..n and earlier events unchanged; every event's "
    "actor allowed and its transition legal; every decision by a person; outbound text within the outward "
    "fields; nothing internal on the supplier view; every file stored with its hash and size and every "
    "inbound sheet's options recorded; every proposal under a mandate flagged, with flags that reproduce."
)
TABLE_HEADER = (
    "| Case | Kind | Supplier behavior | Proposals (in mandate) | Recommended | Decision | By | Final state "
    "| Events | Files | Seconds | Problems |"
)


def render(runs: list[CaseRun], title: str, store: str) -> str:
    violations = sum(len(r.violations) for r in runs)
    errors = sum(1 for r in runs if r.error)
    orders = [r for r in runs if r.kind == "order"]
    decided = [r for r in orders if r.decision]
    by = sorted({r.decided_by for r in decided if r.decided_by}) or ["nobody"]
    in_mandate = sum(1 for r in orders if r.recommended)
    transfers = [r for r in runs if r.kind == "transfer"]
    lines = [
        f"# Case flow simulation: {title}",
        "",
        f"{len(runs)} cases played end to end on `{store}`: {len(orders)} order cases "
        f"({len(decided)} decided, {in_mandate} with a proposal inside the mandate) and {len(transfers)} "
        f"transfer cases (published by the system); {errors} errors, **{violations} invariant violations**. "
        "Every order decision was recorded "
        f"by: {', '.join(by)}.",
        "",
        INVARIANTS,
        "",
        TABLE_HEADER,
        "|---|---|---|---|---|---|---|---|---|---|---|---|",
    ]
    for r in runs:
        problems = "; ".join(r.violations) or (r.error or "")
        cells = [
            r.case_id, r.kind, r.behavior, f"{r.proposals} ({r.in_mandate})", r.recommended or "",
            r.decision or "", r.decided_by or "", r.final_state, str(r.events), str(r.attachments),
            str(r.seconds), problems,
        ]  # fmt: skip
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines) + "\n"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="snm-sim", description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    flow = commands.add_parser("flow", help="play the case flow end to end and check the invariants")
    flow.add_argument("--orders", type=int, default=15)
    flow.add_argument("--transfers", type=int, default=5)
    flow.add_argument("--planning", help="local planning run (default: latest out/planning/*)")
    flow.add_argument(
        "--out", default=str(REPO_ROOT / "out" / "sim" / "flow"), help="fresh local store directory"
    )
    flow.add_argument(
        "--databricks", action="store_true", help="play on the lakehouse case store (appends rows)"
    )
    flow.add_argument("--report", help="write evals/reports/<name>.md")
    args = parser.parse_args(argv)
    if args.databricks:
        from snm_cases.cli import _databricks_service

        service = _databricks_service()
        store = "lakehouse case store"
    else:
        out = Path(args.out)
        if out.exists() and any(out.iterdir()):
            raise SystemExit(f"{out} is not empty; pass a fresh --out directory")
        planning_dir = Path(args.planning) if args.planning else default_planning_dir()
        service = CaseService(LocalCaseStore(out), LocalPlanningData(planning_dir, COST_PARAMETER))
        print(json.dumps(service.sync_engine_run(f"sim-flow-{planning_dir.name}")))
        store = f"local store {out.relative_to(REPO_ROOT) if out.is_relative_to(REPO_ROOT) else out}"
    # The supplier dates its offer by the case store's clock (validity is checked against it).
    runs = run_flow(service, args.orders, args.transfers, service.clock().date())
    text = render(runs, args.report or "ad hoc", store)
    print(text)
    if args.report:
        (REPO_ROOT / "evals" / "reports" / f"{args.report}.md").write_text(text, encoding="utf-8")
    return 1 if any(r.violations or r.error for r in runs) else 0


if __name__ == "__main__":
    raise SystemExit(main())
